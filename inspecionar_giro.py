import pandas as pd
from openpyxl import load_workbook

def main():
    xlsx_path = 'GIRO.xlsx'
    try:
        wb = load_workbook(filename=xlsx_path, read_only=True, data_only=True)
        sheet_names = wb.sheetnames
        print(f"Planilhas encontradas: {sheet_names}")
    except Exception as e:
        print(f"Erro ao abrir workbook: {e}")
        return

    # Ler a primeira planilha com pandas
    try:
        df = pd.read_excel(xlsx_path, sheet_name=0)
    except Exception as e:
        print(f"Erro ao ler a planilha com pandas: {e}")
        return

    # Mostrar informações básicas
    print("\nDimensões da planilha:", df.shape)
    print("\nColunas:")
    for i, c in enumerate(df.columns):
        print(f"  {i}: {c}")

    print("\nAmostra de 5 linhas:")
    with pd.option_context('display.max_columns', None):
        print(df.head(5))

    # Detectar possíveis colunas de item e mês/data
    cols_norm = {c: str(c).strip().lower() for c in df.columns}
    item_candidates = [c for c,n in cols_norm.items() if any(k in n for k in ['item','produto','código','codigo','sku','material','descrição','descricao'])]
    date_candidates = [c for c,n in cols_norm.items() if any(k in n for k in ['data','mês','mes','competência','competencia','emissão','emissao'])]
    value_candidates = [c for c,n in cols_norm.items() if any(k in n for k in ['giro','quantidade','qtd','saida','consumo','movimento'])]

    print("\nPossíveis colunas de item:", item_candidates)
    print("Possíveis colunas de data/mês:", date_candidates)
    print("Possíveis colunas de valor (giro/quantidade):", value_candidates)

    # Estatística: dias distintos com movimento por item
    try:
        df['DATA_FATURAMENTO'] = pd.to_datetime(df['DATA_FATURAMENTO'], errors='coerce')
        df['QUANTIDADE_FATURADA'] = pd.to_numeric(df['QUANTIDADE_FATURADA'], errors='coerce').fillna(0)
        df = df.dropna(subset=['DATA_FATURAMENTO'])
        df['DIA'] = df['DATA_FATURAMENTO'].dt.date
        dias_por_item = df[df['QUANTIDADE_FATURADA'] > 0].groupby(['CODPROD','DESCRICAO','UNIDADE'])['DIA'].nunique()
        print("\nDistribuição de dias com movimento por item:")
        print("  Min:", dias_por_item.min(), " | Max:", dias_por_item.max(), " | Média:", round(dias_por_item.mean(),2))
        print("  Itens com 1 dia:", int((dias_por_item==1).sum()))
        print("  Itens com >=2 dias:", int((dias_por_item>=2).sum()))
        print("\nExemplos (top 10 por dias):")
        print(dias_por_item.sort_values(ascending=False).head(10))
    except Exception as e:
        print(f"\nFalha ao calcular dias por item: {e}")

if __name__ == '__main__':
    main()